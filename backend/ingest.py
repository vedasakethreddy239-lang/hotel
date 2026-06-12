"""Data ingestion + automatic analytics pipeline for LoyaltyShield AI.

Parses CSV / XLSX / JSON uploads, validates rows, stores them with
source="uploaded", then runs the automatic pipeline:
  1. Derive risk signal flags from explicit columns and/or event history.
  2. Recompute every uploaded account's risk score via the shared engine.
  3. Run graph cluster detection (shared device / IP / guest / property-date).
  4. Register properties (with optional coordinates) for the map dashboard.

All dashboards read from the same tables, so uploaded data drives every
view automatically once ingested.
"""

import csv
import io
import json
from collections import defaultdict
from datetime import datetime, timezone

import models
from risk import SIGNALS, compute_risk

NOW_FMT = "%Y-%m-%dT%H:%M:%SZ"

KNOWN_EVENT_TYPES = {
    "login_new_device", "password_changed", "email_changed", "phone_changed",
    "balance_viewed", "guest_booking_created", "high_value_redemption",
    "destination_cluster_detected", "analyst_review",
}

EVENT_DEFAULTS = {
    "login_new_device": ("Medium", 15),
    "password_changed": ("High", 20),
    "email_changed": ("Medium", 15),
    "phone_changed": ("Medium", 15),
    "balance_viewed": ("Low", 0),
    "guest_booking_created": ("Medium", 15),
    "high_value_redemption": ("High", 20),
    "destination_cluster_detected": ("Critical", 25),
    "analyst_review": ("Low", 0),
}

EVENT_TO_SIGNAL = {
    "login_new_device": "new_device",
    "password_changed": "password_changed",
    "email_changed": "email_changed",
    "phone_changed": "phone_changed",
    "guest_booking_created": "guest_booking",
    "high_value_redemption": "high_value_redemption",
    "destination_cluster_detected": "same_property_cluster",
}

PROFILE_CHANGE_EVENTS = ("password_changed", "email_changed", "phone_changed")
REDEMPTION_EVENTS = ("guest_booking_created", "high_value_redemption")

# Column schemas: kind -> (required columns, optional columns)
SCHEMAS = {
    "accounts": (
        {"account_id"},
        {"loyalty_tier", "account_age_days", "points_balance", "country",
         "last_login_region", "trusted_devices", "signals",
         "days_between_change_and_booking", "created_at"},
    ),
    "events": (
        {"account_id", "event_type", "timestamp"},
        {"severity", "risk_delta", "description"},
    ),
    "bookings": (
        {"account_id", "guest_name", "property_name", "booking_date"},
        {"points_used", "suspicious", "latitude", "longitude", "city", "country"},
    ),
    "devices": (
        {"account_id", "device_id"},
        {"ip_address", "suspicious"},
    ),
}

TEMPLATES = {
    "accounts": "account_id,loyalty_tier,account_age_days,points_balance,country,last_login_region,trusted_devices,signals,days_between_change_and_booking,created_at\n"
                "ACC-001,Gold,820,145000,United States,North America,2,new_device;password_changed;guest_booking,1,2024-02-10T08:00:00Z\n"
                "ACC-002,Silver,365,52000,Germany,Europe,3,,,2025-01-05T12:30:00Z\n",
    "events": "account_id,event_type,timestamp,severity,risk_delta,description\n"
              "ACC-001,login_new_device,2026-05-01T09:15:00Z,Medium,15,Login from unrecognized device\n"
              "ACC-001,password_changed,2026-05-01T09:40:00Z,High,20,Password rotated after login\n"
              "ACC-001,guest_booking_created,2026-05-02T11:00:00Z,Medium,15,Guest booking for third party\n",
    "bookings": "account_id,guest_name,property_name,booking_date,points_used,suspicious,latitude,longitude,city,country\n"
                "ACC-001,Jordan Vale,Grand Meridian Dubai,2026-05-03,85000,true,25.2048,55.2708,Dubai,UAE\n"
                "ACC-002,Riley Stone,Palm Oasis Cancun,2026-04-22,30000,false,21.1619,-86.8515,Cancun,Mexico\n",
    "devices": "account_id,device_id,ip_address,suspicious\n"
               "ACC-001,DVC-5521,185.40.12.7,true\n"
               "ACC-002,DVC-1102,73.55.10.20,false\n",
}

VALID_TIERS = {"Bronze", "Silver", "Gold", "Platinum"}
VALID_SEVERITIES = {"Low", "Medium", "High", "Critical"}


def utcnow_iso():
    return datetime.now(timezone.utc).strftime(NOW_FMT)


# ---------------------------------------------------------------- parsing

def parse_file(filename: str, content: bytes):
    """Return (file_type, {kind: [row dicts]}) or raise ValueError."""
    lower = filename.lower()
    if lower.endswith(".csv"):
        return "csv", {None: _parse_csv(content)}
    if lower.endswith((".xlsx", ".xls")):
        return "xlsx", _parse_xlsx(content)
    if lower.endswith(".json"):
        return "json", _parse_json(content)
    raise ValueError("Unsupported file type. Upload .csv, .xlsx or .json")


def _parse_csv(content: bytes):
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("CSV file has no header row")
    return [{(k or "").strip().lower(): (v.strip() if isinstance(v, str) else v)
             for k, v in row.items()} for row in reader]


def _parse_xlsx(content: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
        sheet_rows = []
        for raw in rows[1:]:
            if raw is None or all(c is None for c in raw):
                continue
            sheet_rows.append({header[i]: ("" if raw[i] is None else str(raw[i]).strip())
                               for i in range(min(len(header), len(raw)))})
        if sheet_rows:
            # sheet name may name the kind; else auto-detect later
            key = ws.title.strip().lower()
            out[key if key in SCHEMAS else None] = sheet_rows
    if not out:
        raise ValueError("Workbook contains no data rows")
    return out


def _parse_json(content: bytes):
    data = json.loads(content.decode("utf-8-sig", errors="replace"))
    out = {}
    if isinstance(data, dict):
        for k, v in data.items():
            kk = k.strip().lower()
            if kk in SCHEMAS and isinstance(v, list):
                out[kk] = [_norm_json_row(r) for r in v if isinstance(r, dict)]
        if not out:
            raise ValueError("JSON object must contain one of: accounts, events, bookings, devices (arrays)")
    elif isinstance(data, list):
        out[None] = [_norm_json_row(r) for r in data if isinstance(r, dict)]
    else:
        raise ValueError("JSON must be an array of objects or an object of arrays")
    return out


def _norm_json_row(row: dict):
    return {str(k).strip().lower(): ("" if v is None else (v if isinstance(v, (int, float, bool)) else str(v).strip()))
            for k, v in row.items()}


def detect_kind(rows):
    if not rows:
        return None
    cols = set(rows[0].keys())
    # order matters: most specific first
    for kind in ("events", "bookings", "devices", "accounts"):
        req, _ = SCHEMAS[kind]
        if req.issubset(cols):
            return kind
    return None


# ---------------------------------------------------------------- validation

def _to_int(val, default=None):
    if val in (None, ""):
        return default
    try:
        return int(float(val))
    except (ValueError, TypeError):
        raise ValueError(f"'{val}' is not a number")


def _to_float(val):
    if val in (None, ""):
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        raise ValueError(f"'{val}' is not a number")


def _to_bool(val):
    if isinstance(val, bool):
        return val
    return str(val).strip().lower() in ("true", "1", "yes", "y")


def _parse_ts(val):
    if val in (None, ""):
        raise ValueError("timestamp is required")
    s = str(val).strip().replace(" ", "T", 1)
    for fmt in ("%Y-%m-%dT%H:%M:%SZ", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).strftime(NOW_FMT)
        except ValueError:
            continue
    raise ValueError(f"unrecognized timestamp '{val}' (expected ISO 8601)")


def _parse_date(val):
    s = str(val).strip()[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d").strftime("%Y-%m-%d")
    except ValueError:
        raise ValueError(f"unrecognized date '{val}' (expected YYYY-MM-DD)")


def _parse_signals(val):
    """Accept JSON object, JSON list, or semicolon/comma separated keys."""
    if val in (None, ""):
        return {}
    if isinstance(val, str) and val.lstrip().startswith(("{", "[")):
        parsed = json.loads(val)
        if isinstance(parsed, dict):
            keys = [k for k, v in parsed.items() if v]
        else:
            keys = list(parsed)
    else:
        keys = [k.strip() for k in str(val).replace(",", ";").split(";") if k.strip()]
    bad = [k for k in keys if k not in SIGNALS]
    if bad:
        raise ValueError(f"unknown signal(s): {', '.join(bad)}")
    return {k: True for k in keys}


def validate_rows(kind, rows):
    """Return (clean_rows, errors). Errors: [{row, error}]."""
    req, opt = SCHEMAS[kind]
    clean, errors = [], []
    for idx, row in enumerate(rows, start=2):  # +1 header, 1-based
        try:
            missing = [c for c in req if not str(row.get(c, "")).strip()]
            if missing:
                raise ValueError(f"missing required: {', '.join(missing)}")
            r = {}
            if kind == "accounts":
                r["account_id"] = str(row["account_id"]).strip()
                tier = str(row.get("loyalty_tier", "") or "Bronze").strip().title()
                if tier not in VALID_TIERS:
                    raise ValueError(f"invalid loyalty_tier '{tier}'")
                r["loyalty_tier"] = tier
                r["account_age_days"] = max(0, _to_int(row.get("account_age_days"), 0))
                r["points_balance"] = max(0, _to_int(row.get("points_balance"), 0))
                r["country"] = str(row.get("country", "") or "Unknown")
                r["last_login_region"] = str(row.get("last_login_region", "") or "Unknown")
                r["trusted_devices"] = max(0, _to_int(row.get("trusted_devices"), 1))
                r["signals"] = _parse_signals(row.get("signals"))
                r["days_between_change_and_booking"] = _to_int(row.get("days_between_change_and_booking"), None)
                r["created_at"] = _parse_ts(row.get("created_at")) if str(row.get("created_at", "")).strip() else utcnow_iso()
            elif kind == "events":
                r["account_id"] = str(row["account_id"]).strip()
                et = str(row["event_type"]).strip().lower().replace(" ", "_")
                r["event_type"] = et
                r["timestamp"] = _parse_ts(row["timestamp"])
                d_sev, d_delta = EVENT_DEFAULTS.get(et, ("Low", 0))
                sev = str(row.get("severity", "") or d_sev).strip().title()
                if sev not in VALID_SEVERITIES:
                    raise ValueError(f"invalid severity '{sev}'")
                r["severity"] = sev
                r["risk_delta"] = _to_int(row.get("risk_delta"), d_delta)
                r["description"] = str(row.get("description", "") or f"Uploaded event: {et}")
            elif kind == "bookings":
                r["account_id"] = str(row["account_id"]).strip()
                r["guest_name"] = str(row["guest_name"]).strip()
                r["property_name"] = str(row["property_name"]).strip()
                r["booking_date"] = _parse_date(row["booking_date"])
                r["points_used"] = max(0, _to_int(row.get("points_used"), 0))
                r["suspicious"] = _to_bool(row.get("suspicious", False))
                r["latitude"] = _to_float(row.get("latitude"))
                r["longitude"] = _to_float(row.get("longitude"))
                if r["latitude"] is not None and not -90 <= r["latitude"] <= 90:
                    raise ValueError("latitude out of range")
                if r["longitude"] is not None and not -180 <= r["longitude"] <= 180:
                    raise ValueError("longitude out of range")
                r["city"] = str(row.get("city", "") or "")
                r["country"] = str(row.get("country", "") or "")
            elif kind == "devices":
                r["account_id"] = str(row["account_id"]).strip()
                r["device_id"] = str(row["device_id"]).strip()
                r["ip_address"] = str(row.get("ip_address", "") or "0.0.0.0")
                r["suspicious"] = _to_bool(row.get("suspicious", False))
            clean.append(r)
        except (ValueError, json.JSONDecodeError) as e:
            errors.append({"row": idx, "error": str(e)})
    return clean, errors


# ---------------------------------------------------------------- pipeline

def ingest(db, filename: str, content: bytes):
    """Full ingest: parse -> validate -> store -> pipeline. Returns Dataset."""
    file_type, sheets = parse_file(filename, content)

    accepted = {"accounts": 0, "events": 0, "bookings": 0, "devices": 0}
    all_errors, rows_total = [], 0

    for key, rows in sheets.items():
        rows_total += len(rows)
        kind = key or detect_kind(rows)
        if kind is None:
            cols = ", ".join(sorted(rows[0].keys())) if rows else "none"
            raise ValueError(
                f"Could not detect data type from columns [{cols}]. "
                "Expected accounts (account_id), events (account_id,event_type,timestamp), "
                "bookings (account_id,guest_name,property_name,booking_date) or devices (account_id,device_id)."
            )
        clean, errors = validate_rows(kind, rows)
        for e in errors:
            e["kind"] = kind
        all_errors.extend(errors)
        accepted[kind] += _store(db, kind, clean)

    if sum(accepted.values()) == 0:
        raise ValueError("No valid rows found. " + (f"First error: row {all_errors[0]['row']}: {all_errors[0]['error']}" if all_errors else ""))

    pipeline_result = run_pipeline(db)

    ds = models.Dataset(
        name=filename,
        file_type=file_type,
        kind=", ".join(k for k, v in accepted.items() if v) or "unknown",
        status="completed",
        rows_total=rows_total,
        rows_accepted=sum(accepted.values()),
        rows_rejected=len(all_errors),
        errors=json.dumps(all_errors[:50]),
        summary=json.dumps({"accepted": accepted, "pipeline": pipeline_result}),
        created_at=utcnow_iso(),
    )
    db.add(ds)
    db.commit()
    db.refresh(ds)
    return ds


def _store(db, kind, rows):
    """Insert/upsert validated rows with source='uploaded'."""
    n = 0
    if kind == "accounts":
        for r in rows:
            acc = db.query(models.Account).filter(models.Account.account_id == r["account_id"]).first()
            sig_json = json.dumps(r["signals"])
            if acc:
                acc.loyalty_tier = r["loyalty_tier"]
                acc.account_age_days = r["account_age_days"]
                acc.points_balance = r["points_balance"]
                acc.country = r["country"]
                acc.last_login_region = r["last_login_region"]
                acc.trusted_devices = r["trusted_devices"]
                acc.signals = sig_json
                acc.days_between_change_and_booking = r["days_between_change_and_booking"]
                acc.source = "uploaded"
            else:
                db.add(models.Account(
                    account_id=r["account_id"], loyalty_tier=r["loyalty_tier"],
                    account_age_days=r["account_age_days"], points_balance=r["points_balance"],
                    country=r["country"], last_login_region=r["last_login_region"],
                    trusted_devices=r["trusted_devices"], risk_score=0, risk_status="Low",
                    signals=sig_json, days_between_change_and_booking=r["days_between_change_and_booking"],
                    reviewed=False, source="uploaded", created_at=r["created_at"],
                ))
            n += 1
    elif kind == "events":
        for r in rows:
            db.add(models.Event(source="uploaded", **r))
            n += 1
    elif kind == "bookings":
        for r in rows:
            _ensure_property(db, r)
            db.add(models.Booking(
                account_id=r["account_id"], guest_name=r["guest_name"],
                property_name=r["property_name"], booking_date=r["booking_date"],
                points_used=r["points_used"], suspicious=r["suspicious"], source="uploaded",
            ))
            n += 1
    elif kind == "devices":
        for r in rows:
            db.add(models.AccountDevice(source="uploaded", **r))
            n += 1
    db.flush()
    return n


def _ensure_property(db, r):
    prop = db.query(models.Property).filter(models.Property.name == r["property_name"]).first()
    if prop:
        # backfill coordinates if newly provided
        if prop.latitude is None and r.get("latitude") is not None:
            prop.latitude, prop.longitude = r["latitude"], r["longitude"]
        if not prop.city and r.get("city"):
            prop.city = r["city"]
        if not prop.country and r.get("country"):
            prop.country = r["country"]
    else:
        db.add(models.Property(
            name=r["property_name"], city=r.get("city") or "", country=r.get("country") or "",
            latitude=r.get("latitude"), longitude=r.get("longitude"), source="uploaded",
        ))
        db.flush()


def _ensure_stub_accounts(db):
    """Create minimal accounts for uploaded events/bookings/devices that
    reference unknown account_ids, so every view can link through."""
    known = {a.account_id for a in db.query(models.Account.account_id).all()}
    referenced = set()
    for model in (models.Event, models.Booking, models.AccountDevice):
        for (aid,) in db.query(model.account_id).filter(model.source == "uploaded").distinct().all():
            referenced.add(aid)
    created = 0
    for aid in sorted(referenced - known):
        db.add(models.Account(
            account_id=aid, loyalty_tier="Bronze", account_age_days=0, points_balance=0,
            country="Unknown", last_login_region="Unknown", trusted_devices=1,
            risk_score=0, risk_status="Low", signals="{}", reviewed=False,
            source="uploaded", created_at=utcnow_iso(),
        ))
        created += 1
    db.flush()
    return created


def run_pipeline(db):
    """Recompute risk + rebuild clusters for the uploaded dataset."""
    stub_created = _ensure_stub_accounts(db)

    accounts = db.query(models.Account).filter(models.Account.source == "uploaded").all()
    events_by_acc = defaultdict(list)
    for e in db.query(models.Event).filter(models.Event.source == "uploaded").all():
        events_by_acc[e.account_id].append(e)
    bookings = db.query(models.Booking).filter(models.Booking.source == "uploaded").all()
    devices = db.query(models.AccountDevice).filter(models.AccountDevice.source == "uploaded").all()
    bookings_by_acc = defaultdict(list)
    for b in bookings:
        bookings_by_acc[b.account_id].append(b)

    clusters = _detect_clusters(db, accounts, bookings, devices)
    clustered_accounts = set()
    for c in clusters:
        if c.cluster_type == "same_property_dates":
            clustered_accounts.update(json.loads(c.account_ids))

    rescored = 0
    for acc in accounts:
        flags = json.loads(acc.signals or "{}")
        evts = sorted(events_by_acc.get(acc.account_id, []), key=lambda e: e.timestamp)
        for e in evts:
            sig = EVENT_TO_SIGNAL.get(e.event_type)
            if sig:
                flags[sig] = True
            if e.event_type == "balance_viewed" and acc.account_age_days > 700:
                flags["dormant_balance_check"] = True
        if bookings_by_acc.get(acc.account_id):
            flags["guest_booking"] = True
            if any(b.points_used >= 40000 for b in bookings_by_acc[acc.account_id]):
                flags["high_value_redemption"] = True
        if acc.account_id in clustered_accounts:
            flags["same_property_cluster"] = True
        # protective signals only apply while nothing anomalous is observed;
        # strip auto-added ones once anomalies appear (re-runs stay consistent)
        anomalous = any(flags.get(k) for k in SIGNALS if SIGNALS[k]["weight"] > 0)
        if anomalous:
            flags.pop("trusted_device", None)
            flags.pop("long_clean_history", None)
        else:
            if acc.trusted_devices >= 1:
                flags["trusted_device"] = True
            if acc.account_age_days >= 730:
                flags["long_clean_history"] = True

        days_between = acc.days_between_change_and_booking
        if days_between is None:
            days_between = _derive_days_between(evts, bookings_by_acc.get(acc.account_id, []))
            acc.days_between_change_and_booking = days_between

        result = compute_risk(flags, days_between)
        acc.risk_score = result["risk_score"]
        acc.risk_status = result["risk_status"]
        acc.signals = json.dumps({k: v for k, v in flags.items() if v})
        rescored += 1

    db.commit()
    return {
        "accounts_rescored": rescored,
        "stub_accounts_created": stub_created,
        "clusters_detected": len(clusters),
        "cluster_ids": [c.cluster_id for c in clusters],
    }


def _derive_days_between(events, acc_bookings):
    """Days between first profile-change event and first redemption activity."""
    change_ts = next((e.timestamp for e in events if e.event_type in PROFILE_CHANGE_EVENTS), None)
    if not change_ts:
        return None
    redeem_ts = next((e.timestamp for e in events if e.event_type in REDEMPTION_EVENTS), None)
    if not redeem_ts and acc_bookings:
        first = min(b.booking_date for b in acc_bookings)
        redeem_ts = f"{first}T00:00:00Z"
    if not redeem_ts or redeem_ts < change_ts:
        return None
    d1 = datetime.strptime(change_ts, NOW_FMT)
    d2 = datetime.strptime(redeem_ts, NOW_FMT)
    return max(0, (d2 - d1).days)


def _detect_clusters(db, accounts, bookings, devices):
    """Rebuild uploaded-source clusters from scratch each run."""
    db.query(models.Cluster).filter(models.Cluster.source == "uploaded").delete()
    db.flush()

    clusters = []
    num = 0

    def add(ctype, severity, description, account_ids):
        nonlocal num
        num += 1
        c = models.Cluster(
            cluster_id=f"UCL-{num:03d}", cluster_type=ctype, severity=severity,
            description=description, account_ids=json.dumps(sorted(account_ids)),
            source="uploaded",
        )
        db.add(c)
        clusters.append(c)

    # shared device
    by_device = defaultdict(set)
    for d in devices:
        by_device[d.device_id].add(d.account_id)
    for dev, accs in by_device.items():
        if len(accs) >= 2:
            add("shared_device", "Critical",
                f"Device {dev} used to access {len(accs)} unrelated loyalty accounts.", accs)
            for d in devices:
                if d.device_id == dev:
                    d.suspicious = True

    # shared IP (3+ accounts)
    by_ip = defaultdict(set)
    for d in devices:
        if d.ip_address and d.ip_address != "0.0.0.0":
            by_ip[d.ip_address].add(d.account_id)
    for ip, accs in by_ip.items():
        if len(accs) >= 3:
            add("shared_ip", "Medium",
                f"IP address {ip} associated with logins on {len(accs)} accounts.", accs)
            for d in devices:
                if d.ip_address == ip:
                    d.suspicious = True

    # shared guest name
    by_guest = defaultdict(set)
    for b in bookings:
        by_guest[b.guest_name].add(b.account_id)
    for guest, accs in by_guest.items():
        if len(accs) >= 2:
            add("shared_guest", "High",
                f"Guest '{guest}' appears in bookings across {len(accs)} unrelated accounts.", accs)
            for b in bookings:
                if b.guest_name == guest:
                    b.suspicious = True

    # same property within 3-day window (3+ accounts)
    by_prop = defaultdict(list)
    for b in bookings:
        by_prop[b.property_name].append(b)
    for prop, blist in by_prop.items():
        blist.sort(key=lambda b: b.booking_date)
        i = 0
        while i < len(blist):
            start = datetime.strptime(blist[i].booking_date, "%Y-%m-%d")
            window = [b for b in blist if 0 <= (datetime.strptime(b.booking_date, "%Y-%m-%d") - start).days <= 3]
            accs = {b.account_id for b in window}
            if len(accs) >= 3:
                add("same_property_dates", "High",
                    f"{len(accs)} accounts redeemed at {prop} within a 72-hour window.", accs)
                for b in window:
                    b.suspicious = True
                i += len(window)
            else:
                i += 1

    db.flush()
    return clusters


def clear_uploaded(db):
    """Remove all uploaded data, restoring pure demo mode."""
    counts = {}
    for model in (models.Account, models.Event, models.Booking, models.AccountDevice,
                  models.Cluster, models.Property):
        counts[model.__tablename__] = (
            db.query(model).filter(model.source == "uploaded").delete()
        )
    counts["datasets"] = db.query(models.Dataset).delete()
    db.commit()
    return counts


def active_source(db) -> str:
    """'uploaded' when any uploaded accounts exist, else 'demo'."""
    has = db.query(models.Account.id).filter(models.Account.source == "uploaded").first()
    return "uploaded" if has else "demo"
